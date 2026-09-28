"""Excel export: the price list rebuilt as a real spreadsheet (new prices as
numbers), plus a sheet listing every change."""
from __future__ import annotations

import concurrent.futures
import logging
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import ai, config
from .models import Analysis, PriceItem
from .pipeline import AI_TEXT_EDGE, page_view

log = logging.getLogger(__name__)
if hasattr(pymupdf, "no_recommend_layout"):
    pymupdf.no_recommend_layout()   # silence an advert printed by the table finder

FONT = "Tahoma"
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
PRICE_FILL = PatternFill("solid", fgColor="FFF2CC")
ZEBRA_FILL = PatternFill("solid", fgColor="F7F9FC")


@dataclass
class Cell:
    text: str
    price_id: str = ""


@dataclass
class Table:
    title: str
    direction: str
    headers: list[str]
    rows: list[list[Cell]] = field(default_factory=list)


# ============================================================== extraction ==

def extract_tables(analysis: Analysis) -> tuple[list[Table], str]:
    """Tables of the pages that hold prices. Returns (tables, source) with
    source "ai", "pdf" (PyMuPDF table finder) or "list" (just the prices).
    Every price appears somewhere: prices outside any table are listed last."""
    tables, source = _extract(analysis)
    placed = {c.price_id for t in tables for r in t.rows for c in r if c.price_id}
    rest = [it for it in analysis.items if it.id not in placed]
    if rest and source != "list":
        extra = _list_table(analysis, sorted({it.page for it in rest}), only={it.id for it in rest})
        extra[0].title = "سایر قیمت‌ها"
        tables += extra
    return tables, source


def _extract(analysis: Analysis) -> tuple[list[Table], str]:
    pages = sorted({it.page for it in analysis.items})
    if not pages:
        return [], "list"
    if ai.enabled():
        results: dict[int, list[Table] | None] = {}
        with concurrent.futures.ThreadPoolExecutor(max_workers=config.AI_PARALLEL_PAGES) as pool:
            futs = {pool.submit(_ai_page, analysis, p): p for p in pages}
            for fut in concurrent.futures.as_completed(futs):
                p = futs[fut]
                try:
                    results[p] = fut.result()
                except Exception as exc:  # noqa: BLE001
                    log.warning("table transcription failed on page %d: %s", p + 1, exc)
                    results[p] = None
        if all(results[p] for p in pages):
            return _merge_tables([t for p in pages for t in results[p]]), "ai"
        if any(results[p] for p in pages):
            tables = []
            for p in pages:
                fallback = _pdf_tables(analysis, [p]) if analysis.kind == "pdf" else []
                tables += results[p] or fallback or _list_table(analysis, [p])
            return _merge_tables(tables), "ai"
    if analysis.kind == "pdf":
        tables = _pdf_tables(analysis, pages)
        if tables:
            return _merge_tables(tables), "pdf"
    return _list_table(analysis, pages), "list"


def _ai_page(analysis: Analysis, page: int) -> list[Table]:
    edge = AI_TEXT_EDGE if analysis.pages[page].mode == "text" else ai.MAX_IMAGE_EDGE
    view = page_view(analysis, page, edge)
    items = [it for it in analysis.items if it.page == page]
    prices = [(it.id, it.text, view.boxes[it.id]) for it in items if it.id in view.boxes]
    data = ai.transcribe_tables(view.image, prices)
    valid = {it.id for it in items}
    tables = []
    for t in data.get("tables", []):
        headers = [str(h) for h in t.get("headers", [])]
        rows = []
        for r in t.get("rows", []):
            cells = [Cell(str(c.get("text", "")), c.get("price_id", "") if c.get("price_id") in valid else "")
                     for c in r.get("cells", [])]
            if any(c.text.strip() or c.price_id for c in cells):
                rows.append(cells)
        if headers or rows:
            tables.append(Table(str(t.get("title", "")), t.get("direction", "rtl"), headers, rows))
    missing = valid - {c.price_id for t in tables for r in t.rows for c in r}
    if missing:
        log.info("page %d: %d prices not placed in a table cell", page + 1, len(missing))
    return tables


_ARABIC = re.compile(r"[؀-ۿ]")
_LATIN = re.compile(r"[A-Za-z]")
_ARABIC_WORD = re.compile(r"[؀-ۿ‌]{2,}")
_LTR_RUN = re.compile(r"[0-9A-Za-z۰-۹٠-٩][0-9A-Za-z۰-۹٠-٩.,/:%+\-_٬]*")


def _visual_to_logical(s: str) -> str:
    """Glyphs listed left-to-right (visual order) -> reading order for RTL text."""
    return _LTR_RUN.sub(lambda m: m.group(0)[::-1], s[::-1])


def _fix_order(s: str, page_words: set[str]) -> str:
    """The table finder lists Persian text in visual order on some PDFs and in
    reading order on others; keep whichever version uses real words of the page."""
    if not _ARABIC.search(s):
        return s
    alt = _visual_to_logical(s)

    def hits(x: str) -> int:
        return sum(w in page_words for w in _ARABIC_WORD.findall(x))
    return alt if hits(alt) >= hits(s) else s


def _drop_empty_columns(grid: list[list["Cell"]]) -> list[list["Cell"]]:
    width = max((len(r) for r in grid), default=0)
    keep = [c for c in range(width)
            if any(c < len(r) and (r[c].text.strip() or r[c].price_id) for r in grid)]
    return [[r[c] if c < len(r) else Cell("") for c in keep] for r in grid]


def _pdf_tables(analysis: Analysis, pages: list[int]) -> list[Table]:
    out: list[Table] = []
    doc = pymupdf.open(analysis.source)
    try:
        for pno in pages:
            page = doc[pno]
            text = page.get_text()
            page_words = set(_ARABIC_WORD.findall(text))
            rtl = len(_ARABIC.findall(text)) >= len(_LATIN.findall(text))
            items = [it for it in analysis.items if it.page == pno and it.kind == "text"]
            try:
                found = page.find_tables().tables
            except Exception:  # noqa: BLE001 - table finder is best effort
                found = []
            for tab in found:
                data = tab.extract()
                if len(data) < 2:
                    continue
                grid: list[list[Cell]] = []
                for ri, row in enumerate(data):
                    rects = tab.rows[ri].cells if ri < len(tab.rows) else [None] * len(row)
                    cells = []
                    for ci, txt in enumerate(row):
                        pid = ""
                        rect = rects[ci] if ci < len(rects) else None
                        if rect is not None:
                            r = pymupdf.Rect(rect)
                            for it in items:
                                b = pymupdf.Rect(it.bbox)
                                if r.contains(pymupdf.Point((b.x0 + b.x1) / 2, (b.y0 + b.y1) / 2)):
                                    pid = it.id
                                    break
                        clean = re.sub(r"\s+", " ", (txt or "").replace("\n", " ")).strip()
                        cells.append(Cell(_fix_order(clean, page_words), pid))
                    grid.append(cells[::-1] if rtl else cells)
                grid = _drop_empty_columns(grid)
                if not grid or not grid[0]:
                    continue
                headers = [c.text for c in grid[0]]
                out.append(Table("", "rtl" if rtl else "ltr", headers, grid[1:]))
    finally:
        doc.close()
    return out


def _list_table(analysis: Analysis, pages: list[int], only: set[str] | None = None) -> list[Table]:
    rows = [[Cell(str(it.page + 1)), Cell(it.label or "-"), Cell(it.column or "-"), Cell(it.text, it.id)]
            for it in analysis.items if it.page in pages and (only is None or it.id in only)]
    return [Table("", "rtl", ["صفحه", "شرح", "ستون", "قیمت"], rows)]


def _norm(headers: list[str]) -> tuple[str, ...]:
    return tuple(re.sub(r"\s+", "", h) for h in headers)


def _merge_tables(tables: list[Table]) -> list[Table]:
    """Glue a table that continues on the next page (same header) to the previous one."""
    out: list[Table] = []
    for t in tables:
        if out and _norm(out[-1].headers) == _norm(t.headers) and t.headers:
            out[-1].rows.extend(t.rows)
        else:
            out.append(t)
    return out


# ================================================================= writing ==

def _number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def _num_format(item: PriceItem) -> str:
    return "#,##0" if not item.fmt.decimals else "#,##0." + "0" * item.fmt.decimals


def _cell_value(text: str):
    t = text.strip()
    if re.fullmatch(r"\d{1,6}", t):          # row numbers, quantities
        return int(t)
    return t                                  # codes etc. stay text (no scientific notation)


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "Sheet"
    candidate, k = base, 2
    while candidate in used:
        candidate = f"{base[:25]} {k}"
        k += 1
    used.add(candidate)
    return candidate


def write_workbook(entries: list[tuple[Analysis, dict[str, Decimal], list[Table]]], path: Path,
                   summary: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for analysis, values, tables in entries:
        ws = wb.create_sheet(_sheet_name(Path(analysis.filename).stem, used))
        items = {it.id: it for it in analysis.items}
        rtl = sum(t.direction == "rtl" for t in tables) >= len(tables) / 2
        ws.sheet_view.rightToLeft = rtl
        widths: dict[int, float] = {}
        r = 1
        first_header_row = None
        for t in tables:
            ncol = max([len(t.headers)] + [len(row) for row in t.rows] + [1])
            if t.title:
                ws.cell(r, 1, t.title).font = Font(name=FONT, bold=True, size=13)
                ws.cell(r, 1).fill = TITLE_FILL
                ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
                if ncol > 1:
                    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
                ws.row_dimensions[r].height = 24
                r += 1
            if t.headers:
                for c, h in enumerate(t.headers, 1):
                    cell = ws.cell(r, c, h)
                    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
                    cell.fill = HEAD_FILL
                    cell.border = BORDER
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    widths[c] = max(widths.get(c, 0), len(h))
                ws.row_dimensions[r].height = 30
                first_header_row = first_header_row or r
                r += 1
            for k, row in enumerate(t.rows):
                for c, cell_data in enumerate(row, 1):
                    item = items.get(cell_data.price_id) if cell_data.price_id else None
                    if item is not None:
                        value = values.get(item.id, item.value)
                        cell = ws.cell(r, c, _number(value))
                        cell.number_format = _num_format(item)
                        cell.font = Font(name=FONT, bold=True)
                        cell.fill = PRICE_FILL
                        shown = f"{value:,}"
                    else:
                        cell = ws.cell(r, c, _cell_value(cell_data.text))
                        cell.font = Font(name=FONT)
                        if k % 2:
                            cell.fill = ZEBRA_FILL
                        shown = cell_data.text
                    cell.border = BORDER
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    widths[c] = max(widths.get(c, 0), len(shown))
                r += 1
            r += 1
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(8.0, w * 1.15 + 3), 60.0)
        if first_header_row and len(tables) == 1:
            ws.freeze_panes = ws.cell(first_header_row + 1, 1)
        ws.page_setup.orientation = "landscape" if len(widths) > 6 else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    _changes_sheet(wb, entries, summary)
    wb.properties.title = "لیست قیمت"
    wb.save(path)


def _row_labels(tables: list[Table]) -> dict[str, str]:
    """price id -> 'row number - description' taken from the table row it sits in."""
    out: dict[str, str] = {}
    for t in tables:
        for row in t.rows:
            ids = [c.price_id for c in row if c.price_id]
            if not ids:
                continue
            texts = [c.text.strip() for c in row if not c.price_id and c.text.strip()]
            number = next((x for x in texts if re.fullmatch(r"\d{1,4}", x)), "")
            desc = max((x for x in texts if _ARABIC.search(x) or _LATIN.search(x)), key=len, default="")
            label = " - ".join(x for x in (number, desc) if x)
            for pid in ids:
                out[pid] = label or "-"
    return out


def _changes_sheet(wb: Workbook, entries, summary: str) -> None:
    ws = wb.create_sheet("تغییرات")
    ws.sheet_view.rightToLeft = True
    ws.cell(1, 1, f"خلاصه تغییرات — {summary}").font = Font(name=FONT, bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    ws.cell(1, 1).fill = TITLE_FILL
    headers = ["فایل", "صفحه", "شرح", "ستون", "قیمت قبلی", "قیمت جدید", "تغییر"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(2, c, h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r = 3
    for analysis, values, tables in entries:
        row_text = _row_labels(tables)
        for it in analysis.items:
            if it.id not in values:
                continue
            new = values[it.id]
            row = [analysis.filename, it.page + 1, it.label or row_text.get(it.id, "-"), it.column or "-",
                   _number(it.value), _number(new), float((new - it.value) / it.value) if it.value else 0.0]
            for c, v in enumerate(row, 1):
                cell = ws.cell(r, c, v)
                cell.font = Font(name=FONT)
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(r, 5).number_format = _num_format(it)
            ws.cell(r, 6).number_format = _num_format(it)
            ws.cell(r, 7).number_format = "0.00%"
            r += 1
    for c, w in zip(range(1, 8), (22, 7, 40, 18, 16, 16, 10)):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A3"
